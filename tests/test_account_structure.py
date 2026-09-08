from __future__ import annotations

import asyncio
import io
import sqlite3
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from fastapi.testclient import TestClient
from fastapi import UploadFile

import test_quality as baseline
from test_quality import app, quality
from qcc import access
from qcc.account_schema import sampling_config, statements


class AccountStructureTests(unittest.TestCase):
    setUp = baseline.QualityCommandCenterTests.setUp
    create_scorecard = baseline.QualityCommandCenterTests.create_scorecard

    def account(self, name):
        aid = app.create_account(app.AccountIn(name=name), user="admin")["id"]
        pid = quality.list_processes(account_id=aid, user="admin")[0]["id"]
        return aid, pid

    def user(self, name, grants=()):
        quality.save_admin_user(quality.UserIn(username=name, display_name=name, password="test-password-123",
            account_roles=[quality.AccountRoleIn(account_id=a, roles=list(r)) for a, r in grants], must_change_password=False), user="admin")

    def case(self, pid, name, card=None, status="REVIEWED"):
        with app.db() as con:
            con.execute("""INSERT INTO audit_cases(audit_id,process_id,scorecard_version_id,status,weighted_score,passed,
                opportunities,defect_count,created_at,reviewed_at,created_by) VALUES(?,?,?,?,100,1,10,0,?,?,'admin')""",
                (name, pid, card, status, app.iso_now(), app.iso_now() if status == "REVIEWED" else None))

    def upload(self, name="sample.csv", content=b"Case ID,Associate\nA1,Alice\nA2,Alice\nA3,Alice\n", user="admin"):
        return asyncio.run(app.upload_workbook(UploadFile(filename=name, file=io.BytesIO(content)), user=user))

    def payload(self, aid, pid, **kwargs):
        return app.RunIn(account_id=aid, process_id=pid, sheet_name="CSV", header_row=1,
            identifier_column="Case ID", associate_column="Associate", requested_count=1, **kwargs)

    def test_account_roles_scope_lists_records_aggregates_and_exports(self):
        a, p = self.account("Account A"); b, q = self.account("Account B")
        self.user("auditor", [(a, ["QA Auditor"])])
        self.user("pending")
        ca = self.create_scorecard(p); cb = self.create_scorecard(q)
        self.case(p, "A-CASE", ca); self.case(q, "B-CASE", cb)
        capa_a = quality.create_capa(quality.CapaIn(process_id=p, title="Account A issue"), user="admin")["capa_id"]
        capa_b = quality.create_capa(quality.CapaIn(process_id=q, title="Account B issue"), user="admin")["capa_id"]
        client = TestClient(app.app)
        app.app.dependency_overrides[app.current_user] = lambda: "auditor"
        try:
            for route in ["/api/accounts", "/api/admin/processes", "/api/scorecards", "/api/audits", "/api/capas"]:
                response = client.get(route)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertNotIn("Account B", response.text)
            self.assertEqual(client.get('/api/analytics/summary').json()["metrics"]["units"], 1)
            self.assertEqual(len(client.get('/api/analytics/summary').json()["capas"]), 1)
            for route in [f"/api/audits/B-CASE", f"/api/capas/{capa_b}", f"/api/scorecards?account_id={b}",
                          f"/api/analytics/summary?process_id={q}", f"/api/accounts/{b}/config",
                          f"/api/admin/processes/{q}/sampling-controls", f"/api/reports/quality.xlsx?account_id={b}"]:
                self.assertEqual(client.get(route).status_code, 403, route)
            report = client.get('/api/reports/quality.xlsx')
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(report.content))
            values = str([list(ws.values) for ws in wb.worksheets])
            self.assertIn("A-CASE", values); self.assertNotIn("B-CASE", values)
            self.assertEqual(client.get('/api/admin/users').status_code, 403)
            self.assertEqual(client.get(f'/api/capas/{capa_a}').status_code, 200)
        finally:
            app.app.dependency_overrides.clear()
        self.assertEqual(app.list_accounts(user="pending"), [])
        self.assertEqual(quality.analytics_payload(None, None, None, None, "pending")["metrics"]["units"], 0)
        self.assertTrue(quality.user_context("pending")["assignment_required"])
        self.assertEqual(len(app.list_accounts(user="admin")), 2)

    def test_roles_apply_to_record_account_and_revocation_is_immediate(self):
        a, p = self.account("Account A"); b, q = self.account("Account B")
        self.user("mixed", [(a, ["QA Auditor"]), (b, ["QA Reviewer"])])
        ca = self.create_scorecard(p); cb = self.create_scorecard(q)
        self.case(p, "A-SUBMITTED", ca, "SUBMITTED"); self.case(q, "B-SUBMITTED", cb, "SUBMITTED")
        with self.assertRaises(app.HTTPException) as denied:
            quality.review_audit("A-SUBMITTED", quality.ReviewIn(decision="APPROVE"), user="mixed")
        self.assertEqual(denied.exception.status_code, 403)
        quality.review_audit("B-SUBMITTED", quality.ReviewIn(decision="APPROVE"), user="mixed")
        self.case(q, "B-DRAFT", cb, "UNASSIGNED")
        with self.assertRaises(app.HTTPException):
            quality.save_audit("B-DRAFT", quality.AuditSaveIn(), user="mixed")
        quality.save_admin_user(quality.UserIn(username="mixed", display_name="Mixed", account_roles=[]), user="admin")
        self.assertEqual(quality.list_audits(user="mixed"), [])
        self.assertEqual(access.ACCOUNT_ROLES, {"QA Auditor", "QA Reviewer", "Operations Manager"})
        with self.assertRaises(app.HTTPException):
            quality.save_admin_user(quality.UserIn(username="mixed", display_name="Mixed", roles=["QA Reviewer"]), user="admin")
        with self.assertRaises(app.HTTPException):
            quality.save_admin_user(quality.UserIn(username="admin", display_name="Admin", roles=[]), user="admin")

    def test_sampling_runs_inventory_and_staging_are_scoped(self):
        a, p = self.account("Account A"); b, q = self.account("Account B")
        self.user("auditor", [(a, ["QA Auditor"])])
        self.user("other", [(a, ["QA Auditor"])])
        uploads = [self.upload(user="auditor"), self.upload()]
        run_a = app.generate_run(uploads[0]["upload_id"], self.payload(a, p), user="auditor")["run_id"]
        run_b = app.generate_run(uploads[1]["upload_id"], self.payload(b, q), user="admin")["run_id"]
        self.assertEqual([r["run_id"] for r in app.list_runs(user="auditor")], [run_a])
        self.assertEqual([r["run_id"] for r in app.inventory(user="auditor")], [run_a])
        for function in [app.get_run, app.export_run]:
            with self.assertRaises(app.HTTPException): function(run_b, user="auditor")
        owned = self.upload(user="auditor")
        with self.assertRaises(app.HTTPException):
            app.inspect_upload(owned["upload_id"], app.InspectIn(sheet_name="CSV", header_row=1), user="other")
        with self.assertRaises(app.HTTPException):
            app.preview_sample(owned["upload_id"], self.payload(b, q), user="auditor")
        with self.assertRaises(app.HTTPException):
            app.preview_sample(owned["upload_id"], self.payload(a, q), user="admin")

    def test_process_controls_and_history_do_not_cross_processes(self):
        a, p = self.account("Account A")
        q = quality.create_process(quality.ProcessIn(account_id=a, name="Second process", process_type="back_office"), user="admin")["id"]
        for pid in [p, q]:
            quality.save_sampling_controls(pid, app.ConfigIn(coverage_enabled=True, audits_per_associate=1), user="admin")
        one = self.upload(); app.generate_run(one["upload_id"], self.payload(a, p), user="admin")
        two = self.upload()
        first = app.preview_sample(two["upload_id"], self.payload(a, p), user="admin")
        second = app.preview_sample(two["upload_id"], self.payload(a, q), user="admin")
        self.assertEqual(first["coverage"]["total_required"], 0)
        self.assertEqual(first["validation"]["previously_sampled"], 1)
        self.assertEqual(second["coverage"]["total_required"], 1)
        self.assertEqual(second["validation"]["previously_sampled"], 0)
        quality.save_sampling_controls(q, app.ConfigIn(audits_per_associate=7, case_insensitive_ids=False), user="admin")
        self.assertEqual(quality.get_sampling_controls(p, user="admin")["audits_per_associate"], 1)
        self.assertEqual(quality.get_sampling_controls(q, user="admin")["audits_per_associate"], 7)
        self.assertEqual(quality.get_sampling_controls(p, user="admin")["identifier_column_default"], "Case ID")
        self.assertIsNone(quality.get_sampling_controls(q, user="admin")["identifier_column_default"])
        with self.assertRaises(app.HTTPException) as retired:
            app.save_config(a, app.ConfigIn(), user="admin")
        self.assertEqual(retired.exception.status_code, 410)

    def test_import_owner_account_and_archive_checks(self):
        a, p = self.account("Account A"); b, q = self.account("Account B")
        self.user("reviewer", [(a, ["QA Reviewer"])])
        self.user("colleague", [(a, ["QA Reviewer"])])
        data = b"Case ID,Defects,Opportunities\nIMPORT-1,,10\n"
        preview = asyncio.run(quality.preview_result_import(UploadFile(filename="results.csv", file=io.BytesIO(data)), user="reviewer"))
        iid = preview["import_id"]
        mapping = {"case_id": "Case ID", "defects": "Defects", "opportunities": "Opportunities"}
        with self.assertRaises(app.HTTPException): quality.import_errors(iid, user="colleague")
        with self.assertRaises(app.HTTPException):
            quality.commit_result_import(iid, quality.ImportCommitIn(process_id=p, mapping=mapping), user="colleague")
        with self.assertRaises(app.HTTPException):
            quality.commit_result_import(iid, quality.ImportCommitIn(process_id=q, mapping=mapping), user="reviewer")
        foreign_card = self.create_scorecard(q)
        with self.assertRaises(app.HTTPException):
            quality.commit_result_import(iid, quality.ImportCommitIn(process_id=p, scorecard_version_id=foreign_card, mapping=mapping), user="admin")
        quality.archive_process(p, user="admin")
        with self.assertRaises(app.HTTPException):
            quality.commit_result_import(iid, quality.ImportCommitIn(process_id=p, mapping=mapping), user="reviewer")
        with app.db() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM audit_cases").fetchone()[0], 0)

    def test_assignees_and_cross_account_child_ids_are_validated(self):
        a, p = self.account("Account A"); b, q = self.account("Account B")
        self.user("auditor_a", [(a, ["QA Auditor"])])
        self.user("auditor_b", [(b, ["QA Auditor"])])
        ca = self.create_scorecard(p); cb = self.create_scorecard(q)
        self.case(p, "LOCAL-DRAFT", ca, "UNASSIGNED")
        self.assertEqual({u["username"] for u in quality.account_assignees(a, user="admin")}, {"admin", "auditor_a"})
        with self.assertRaises(app.HTTPException):
            quality.save_audit("LOCAL-DRAFT", quality.AuditSaveIn(assigned_to="auditor_b"), user="admin")
        with app.db() as con:
            foreign_item = con.execute("SELECT id FROM scorecard_items WHERE scorecard_version_id=?", (cb,)).fetchone()[0]
        with self.assertRaises(app.HTTPException):
            quality.save_audit("LOCAL-DRAFT", quality.AuditSaveIn(defect_item_ids=[foreign_item]), user="admin")
        with self.assertRaises(app.HTTPException):
            quality.create_capa(quality.CapaIn(process_id=q, audit_id="LOCAL-DRAFT", title="Wrong process"), user="admin")
        with self.assertRaises(app.HTTPException):
            quality.update_process(p, quality.ProcessIn(account_id=b, name="Moved", process_type="back_office"), user="admin")

    def test_archive_preserves_work_and_restore_requires_active_account(self):
        a, p = self.account("Account A"); card = self.create_scorecard(p)
        self.case(p, "IN-FLIGHT", card, "UNASSIGNED")
        draft = quality.clone_scorecard(card, user="admin")["id"]
        quality.archive_process(p, user="admin")
        self.assertEqual(quality.list_processes(user="admin"), [])
        self.assertEqual(len(quality.list_processes(user="admin", include_archived=True)), 1)
        for action in [lambda: quality.create_scorecard(quality.ScorecardIn(process_id=p, name="Blocked"), user="admin"),
                       lambda: quality.publish_scorecard(draft, user="admin"),
                       lambda: quality.create_capa(quality.CapaIn(process_id=p, title="New issue"), user="admin"),
                       lambda: quality.save_sampling_controls(p, app.ConfigIn(), user="admin")]:
            with self.assertRaises(app.HTTPException): action()
        up = self.upload()
        with self.assertRaises(app.HTTPException): app.generate_run(up["upload_id"], self.payload(a, p), user="admin")
        item = quality.list_scorecards(process_id=p, user="admin")[1]["items"][0]
        quality.save_audit("IN-FLIGHT", quality.AuditSaveIn(defect_item_ids=[item["id"]]), user="admin")
        quality.submit_audit("IN-FLIGHT", user="admin")
        result = quality.review_audit("IN-FLIGHT", quality.ReviewIn(decision="APPROVE"), user="admin")
        self.assertIsNotNone(result["capa_id"])
        quality.update_capa(result["capa_id"], quality.CapaUpdateIn(containment="Existing work continues"), user="admin")
        quality.restore_process(p, user="admin")
        self.assertEqual(len(quality.list_processes(user="admin")), 1)
        quality.archive_process(p, user="admin"); app.archive_account(a, user="admin")
        with self.assertRaises(app.HTTPException): quality.restore_process(p, user="admin")
        self.assertEqual(quality.get_audit("IN-FLIGHT", user="admin")["status"], "REVIEWED")

    def test_clone_edit_reorder_publish_preserves_original_audit(self):
        a, p = self.account("Account A"); original = self.create_scorecard(p)
        self.case(p, "OLD-RESULT", original)
        original_items = quality.get_audit("OLD-RESULT", user="admin")["items"]
        draft = quality.clone_scorecard(original, user="admin")["id"]
        cloned = next(c for c in quality.list_scorecards(user="admin") if c["id"] == draft)
        self.assertTrue(set(i["id"] for i in original_items).isdisjoint(i["id"] for i in cloned["items"]))
        quality.update_scorecard(draft, quality.ScorecardIn(process_id=p, name="Updated", passing_score=90), user="admin")
        item_id = cloned["items"][0]["id"]
        quality.update_scorecard_item(draft, item_id, quality.ScorecardItemIn(item_type="defect", name="Changed", critical=False), user="admin")
        extra = quality.add_scorecard_item(draft, quality.ScorecardItemIn(item_type="defect", name="Extra"), user="admin")["id"]
        quality.reorder_scorecard(draft, quality.ScorecardOrderIn(item_ids=[extra, item_id]), user="admin")
        quality.delete_scorecard_item(draft, extra, user="admin")
        quality.publish_scorecard(draft, user="admin")
        old = quality.get_audit("OLD-RESULT", user="admin")
        self.assertEqual(old["scorecard_version_id"], original); self.assertEqual(old["weighted_score"], 100)
        self.assertEqual(old["items"], original_items)
        with self.assertRaises(app.HTTPException): quality.delete_scorecard_item(original, original_items[0]["id"], user="admin")
        with self.assertRaises(app.HTTPException): quality.update_scorecard_item(draft, item_id, quality.ScorecardItemIn(item_type="defect", name="No"), user="admin")
        with ThreadPoolExecutor(max_workers=2) as pool:
            versions = list(pool.map(lambda _: quality.clone_scorecard(draft, user="admin")["version"], range(2)))
        self.assertEqual(sorted(versions), [3, 4])

    def test_migration_preserves_controls_and_removes_global_non_admin_grants(self):
        # Exercise the actual migration on a pre-v3 database without modifying app data.
        con = sqlite3.connect(":memory:"); con.row_factory = sqlite3.Row
        con.executescript("""CREATE TABLE users(username TEXT PRIMARY KEY);
            CREATE TABLE roles(name TEXT PRIMARY KEY);
            CREATE TABLE accounts(id INTEGER PRIMARY KEY);
            CREATE TABLE user_roles(username TEXT,role_name TEXT);
            CREATE TABLE uploads(upload_id TEXT);
            CREATE TABLE processes(id INTEGER PRIMARY KEY,account_id INTEGER,updated_at TEXT);
            CREATE TABLE account_config(account_id INTEGER,coverage_enabled INTEGER,coverage_period TEXT,audits_per_associate INTEGER,
            identifier_column_default TEXT,associate_column_default TEXT,exclude_previously_sampled INTEGER,case_insensitive_ids INTEGER);
            INSERT INTO users VALUES('admin'),('legacy');
            INSERT INTO roles VALUES('Administrator'),('QA Auditor');
            INSERT INTO user_roles VALUES('admin','Administrator'),('legacy','QA Auditor');
            INSERT INTO accounts VALUES(1);
            INSERT INTO processes VALUES(1,1,'2026-01-01'),(2,1,'2026-01-01');
            INSERT INTO account_config VALUES(1,1,'month',8,'ID','Agent',0,0);""")
        for statement in statements(): con.execute(statement)
        self.assertEqual([tuple(r) for r in con.execute('SELECT * FROM user_roles')], [('admin', 'Administrator')])
        self.assertEqual([tuple(r) for r in con.execute('SELECT * FROM legacy_user_roles')], [('legacy', 'QA Auditor')])
        self.assertEqual(sampling_config(con, 1)["audits_per_associate"], 8)
        self.assertEqual(sampling_config(con, 2)["identifier_column_default"], "ID")
        con.close()
        a, p = self.account("Real account")
        self.user("scoped", [(a, ["QA Auditor"])])
        quality.save_sampling_controls(p, app.ConfigIn(audits_per_associate=9), user="admin")
        quality.init_quality_db(); quality.init_quality_db()
        self.assertEqual(quality.get_sampling_controls(p, user="admin")["audits_per_associate"], 9)
        self.assertEqual(quality.user_context("scoped")["account_roles"][0]["account_id"], a)


if __name__ == '__main__':
    unittest.main()
